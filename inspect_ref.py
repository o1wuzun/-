from docx import Document
import json
p=r'C:\Users\ASUS\Downloads\201909280107-郭梦杰-临大校园微快递-赵德玉-开题报告.docx'; d=Document(p); out={'sections':[],'paras':[],'tables':[]}
for s in d.sections: out['sections'].append({'page_width':s.page_width.inches,'page_height':s.page_height.inches,'left':s.left_margin.inches,'right':s.right_margin.inches,'top':s.top_margin.inches,'bottom':s.bottom_margin.inches,'header':s.header_distance.inches,'footer':s.footer_distance.inches})
for i,p0 in enumerate(d.paragraphs):
 pf=p0.paragraph_format; out['paras'].append({'i':i,'text':p0.text,'style':p0.style.name,'align':str(p0.alignment),'before':pf.space_before.pt if pf.space_before else None,'after':pf.space_after.pt if pf.space_after else None,'line':pf.line_spacing,'runs':[{'text':r.text,'font':r.font.name,'size':r.font.size.pt if r.font.size else None,'bold':r.bold} for r in p0.runs]})
for ti,t in enumerate(d.tables):
 td={'rows':len(t.rows),'cols':len(t.columns),'cells':[]}
 for ri,row in enumerate(t.rows):
  seen=set()
  for ci,c in enumerate(row.cells):
   key=id(c._tc)
   if key in seen: continue
   seen.add(key)
   td['cells'].append({'ri':ri,'ci':ci,'text':c.text,'width':c.width.inches if c.width else None,'paragraphs':[{'style':p0.style.name,'align':str(p0.alignment),'before':p0.paragraph_format.space_before.pt if p0.paragraph_format.space_before else None,'after':p0.paragraph_format.space_after.pt if p0.paragraph_format.space_after else None,'line':p0.paragraph_format.line_spacing,'runs':[{'text':r.text,'size':r.font.size.pt if r.font.size else None,'bold':r.bold,'font':r.font.name} for r in p0.runs]} for p0 in c.paragraphs]})
 out['tables'].append(td)
open(r'F:\社区\ref_inspect.json','w',encoding='utf-8').write(json.dumps(out,ensure_ascii=False,indent=2))
